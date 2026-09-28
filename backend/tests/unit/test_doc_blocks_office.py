"""PowerPoint, Word y Excel llegan a los consejeros como TEXTO extraído."""
import base64
import io

from app.services.ai.doc_blocks import build_doc_blocks, classify_docs


def _pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Resultados Q3"
    s.placeholders[1].text = "Ventas +12%"
    b = io.BytesIO(); prs.save(b)
    return b.getvalue()


def test_pptx_es_legible_y_se_adjunta_como_texto():
    raw = _pptx()
    legibles, ilegibles = classify_docs([{"filename": "junta.pptx"}, {"filename": "viejo.xls"}])
    assert [d["filename"] for d in legibles] == ["junta.pptx"] and legibles[0]["kind"] == "text"
    assert [d["filename"] for d in ilegibles] == ["viejo.xls"]
    blocks = build_doc_blocks([{**legibles[0], "data": base64.b64encode(raw).decode(), "label": "Documento «junta.pptx»"}])
    assert blocks[0]["type"] == "text" and "Resultados Q3" in blocks[0]["text"] and "Ventas +12%" in blocks[0]["text"]


def test_office_danado_no_revienta():
    legibles, _ = classify_docs([{"filename": "roto.docx"}])
    blocks = build_doc_blocks([{**legibles[0], "data": base64.b64encode(b"no soy docx").decode(), "label": "Doc"}])
    assert "no se pudo extraer" in blocks[0]["text"]
