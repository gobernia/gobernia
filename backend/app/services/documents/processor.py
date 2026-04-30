"""
Procesador de documentos — extrae texto y analiza con Claude.
Corre dentro del worker de Celery (no en el proceso FastAPI).
"""
import anthropic

from app.core.config import settings

_ANALYSIS_PROMPT = """Eres un analista de gobierno corporativo. Analiza el siguiente documento
empresarial y extrae los insights más relevantes para el consejo de administración.

Tipo de documento: {doc_type}
Empresa: {company_context}

Documento:
{text}

Responde en JSON con este formato exacto:
{{
  "summary": "Resumen ejecutivo en 2-3 oraciones",
  "key_findings": ["hallazgo 1", "hallazgo 2", "hallazgo 3"],
  "risks": ["riesgo identificado 1", "riesgo identificado 2"],
  "recommendations": ["recomendación 1", "recomendación 2"]
}}"""


def extract_text_from_content(content: bytes, content_type: str) -> str:
    """Extrae texto según el tipo de archivo. Devuelve texto plano."""
    if content_type == "application/pdf":
        return _extract_pdf(content)
    if content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return _extract_docx(content)
    if content_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ):
        return _extract_xlsx(content)
    # Imágenes: Claude Vision lo procesará directamente en el futuro
    return "[Imagen — análisis visual pendiente]"


def _extract_pdf(content: bytes) -> str:
    try:
        import pdfplumber, io
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages[:20])
    except ImportError:
        return "[PDF — instalar pdfplumber para extracción de texto]"


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document as DocxDocument
        import io
        doc = DocxDocument(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return "[DOCX — instalar python-docx para extracción de texto]"


def _extract_xlsx(content: bytes) -> str:
    try:
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets[:3]:
            for row in sheet.iter_rows(max_row=100, values_only=True):
                line = "\t".join(str(c) for c in row if c is not None)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    except ImportError:
        return "[XLSX — instalar openpyxl para extracción de texto]"


def analyze_document_with_claude(
    text: str,
    doc_type: str,
    company_context: str,
) -> dict:
    """Llama a Claude y retorna el análisis estructurado."""
    if not settings.ANTHROPIC_API_KEY or len(text.strip()) < 50:
        return {
            "summary": "Documento recibido — análisis pendiente de procesamiento.",
            "key_findings": [],
            "risks": [],
            "recommendations": [],
        }

    client = anthropic.Anthropic(api_key=settings.ANTHROPIC_API_KEY)
    prompt = _ANALYSIS_PROMPT.format(
        doc_type=doc_type,
        company_context=company_context,
        text=text[:12_000],   # ~3k tokens, deja margen para el modelo
    )
    message = client.messages.create(
        model=settings.AI_MODEL,
        max_tokens=1024,
        messages=[{"role": "user", "content": prompt}],
    )
    import json, re
    raw = message.content[0].text
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if match:
        return json.loads(match.group())
    return {"summary": raw, "key_findings": [], "risks": [], "recommendations": []}
