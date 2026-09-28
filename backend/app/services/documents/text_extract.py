"""
Texto de los formatos de Office que Claude NO lee de forma nativa (PowerPoint, Word, Excel).
PDF e imágenes no pasan por aquí: viajan como bloques document/image.
"""
import io

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Tope de caracteres por documento (~20k tokens): un Excel de miles de filas no debe
# comerse el contexto del consejero.
MAX_CHARS = 80_000


def _pptx(raw: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    out: list[str] = []
    for n, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t:
                    lines.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    celdas = [c.text.strip() for c in row.cells]
                    if any(celdas):
                        lines.append(" | ".join(celdas))
        if slide.has_notes_slide:
            notas = slide.notes_slide.notes_text_frame.text.strip()
            if notas:
                lines.append(f"(Notas del orador: {notas})")
        if lines:
            out.append(f"--- Lámina {n} ---\n" + "\n".join(lines))
    return "\n\n".join(out)


def _docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            celdas = [c.text.strip() for c in row.cells]
            if any(celdas):
                lines.append(" | ".join(celdas))
    return "\n".join(lines)


def _xlsx(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out: list[str] = []
    for sheet in wb.worksheets[:10]:
        filas = []
        for row in sheet.iter_rows(max_row=500, values_only=True):
            celdas = ["" if c is None else str(c) for c in row]
            if any(celdas):
                filas.append(" | ".join(celdas).rstrip(" |"))
        if filas:
            out.append(f"--- Hoja «{sheet.title}» ---\n" + "\n".join(filas))
    return "\n\n".join(out)


_EXTRACTORES = {PPTX: _pptx, DOCX: _docx, XLSX: _xlsx}


def extraer_texto(raw: bytes, media_type: str) -> str:
    """Texto plano del documento. Vacío si el formato no aplica o el archivo está dañado."""
    fn = _EXTRACTORES.get(media_type)
    if fn is None:
        return ""
    try:
        texto = fn(raw).strip()
    except Exception:
        return ""
    if len(texto) > MAX_CHARS:
        texto = texto[:MAX_CHARS] + "\n[… documento recortado por longitud]"
    return texto
